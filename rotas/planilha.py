"""Importação do Excel, preservando as marcações de CEASA."""

from __future__ import annotations

from dataclasses import asdict, dataclass
from io import BytesIO
from pathlib import Path
import re
import unicodedata

from openpyxl import load_workbook
from openpyxl.styles.colors import COLOR_INDEX


def normalizar(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    return " ".join("".join(c for c in text if not unicodedata.combining(c)).upper().split())


@dataclass(frozen=True)
class Cliente:
    id: str
    nome: str
    cnpj: str
    cidade: str
    vendedor: str
    endereco: str
    segmento: str
    marca: str
    ceasa: bool
    linha: int

    def registro(self) -> dict:
        return asdict(self)


def amarelo(cell) -> bool:
    if cell.fill.patternType != "solid":
        return False
    color = cell.fill.fgColor
    if color.type == "rgb":
        rgb = color.rgb[-6:]
    elif color.type == "indexed" and 0 <= color.indexed < len(COLOR_INDEX):
        rgb = COLOR_INDEX[color.indexed][-6:]
    else:
        return False
    r, g, b = (int(rgb[i:i + 2], 16) for i in (0, 2, 4))
    return r >= 200 and g >= 180 and b < min(r, g) * 0.8


def ler_clientes(source: str | Path | bytes) -> list[Cliente]:
    workbook = load_workbook(BytesIO(source) if isinstance(source, bytes) else source, data_only=True)
    try:
        sheet = workbook["Clientes"] if "Clientes" in workbook.sheetnames else workbook.active
        columns = {normalizar(c.value): c.column for c in sheet[1] if c.value}
        required = {"NOME FANTASIA", "CNPJ", "CIDADE", "ENDERECO", "SEGMENTO"}
        missing = required - columns.keys()
        if missing:
            raise ValueError("Colunas ausentes: " + ", ".join(sorted(missing)))
        clients = []
        seen = set()
        for row in sheet.iter_rows(min_row=2):
            if not any(c.value is not None for c in row):
                continue
            def value(name):
                col = columns.get(name)
                return str(row[col - 1].value or "").strip() if col else ""
            nome = value("NOME FANTASIA")
            cnpj = value("CNPJ")
            if not nome:
                raise ValueError(f"Linha {row[0].row}: nome da loja ausente.")
            identifier = re.sub(r"\D", "", cnpj) or f"linha-{row[0].row}"
            if identifier in seen:
                raise ValueError(f"CNPJ repetido na linha {row[0].row}: {cnpj}. Revise a planilha.")
            seen.add(identifier)
            name = normalizar(nome)
            marca = next((m for m in ("BYD", "FIAT", "JEEP", "HONDA", "NISSAN") if re.search(rf"\b{m}\b", name)), "Outras")
            clients.append(Cliente(identifier, nome, cnpj, value("CIDADE"), value("VENDEDOR"),
                                   value("ENDERECO"), value("SEGMENTO") or "Sem segmento",
                                   marca, any(amarelo(c) for c in row), row[0].row))
        return clients
    finally:
        workbook.close()


def filtrar(clientes: list[Cliente], segmentos: list[str], incluir_ceasa: bool = False,
            marcas: list[str] | None = None) -> list[Cliente]:
    return [c for c in clientes if c.segmento in segmentos and (incluir_ceasa or not c.ceasa)
            and (marcas is None or c.marca in marcas)]
