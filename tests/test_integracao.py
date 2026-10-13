from io import BytesIO

import pytest
from openpyxl import load_workbook

from rotas.dados import ROOT, carregar, selecionar
from rotas.otimizador import Regras, agenda, otimizar
from rotas.planilha import ler_clientes, filtrar
from rotas.relatorios import tabelas, excel, html


def test_planilha_real_e_marcas():
    clientes = ler_clientes(ROOT / 'Endereço Clientes.xlsx')
    assert len(clientes) == 19
    assert '07950180001193' not in {c.id for c in clientes}
    assert sum(c.ceasa for c in clientes) == 5
    assert len(filtrar(clientes, ['DAHRUJ', 'DEALER'])) == 14
    assert len(filtrar(clientes, ['DAHRUJ'])) == 12
    assert len(filtrar(clientes, ['DEALER'], marcas=['BYD'])) == 2
    assert filtrar(clientes, []) == []
    assert filtrar(clientes, ['DAHRUJ'], marcas=[]) == []


@pytest.mark.parametrize('ceasa,total,dias_esperados', [(False, 14, 4), (True, 19, 5)])
def test_todas_as_lojas_sem_omissao_com_retorno(ceasa, total, dias_esperados):
    catalogo, matriz = carregar()
    assert '07950180001193' not in matriz['ids']
    assert '07950180001193' not in {p['id'] for p in catalogo['pontos']}
    clientes = filtrar(ler_clientes(ROOT / 'Endereço Clientes.xlsx'), ['DAHRUJ', 'DEALER'], ceasa)
    pontos, d, t = selecionar(clientes, catalogo, matriz)
    regras = Regras()
    dias = otimizar(d, t, regras)
    assert len(dias) == dias_esperados
    assert sorted(j for dia in dias for j in dia.ordem) == list(range(1, total + 1))
    for dia in dias:
        path = (0,) + dia.ordem + (0,)
        assert dia.km == pytest.approx(sum(d[a][b] for a, b in zip(path, path[1:])))
        assert agenda(dia.ordem, t, regras)[-1]['fim'] == pytest.approx(dia.retorno)
        assert dia.retorno <= regras.fim
    resumo, trechos, agenda_df = tabelas(dias, pontos, d, t, regras, 12, 6.34)
    assert trechos['Distância (km)'].sum() == pytest.approx(resumo['Distância (km)'].sum())
    assert trechos['Custo +10% (R$)'].sum() == pytest.approx(resumo['Custo +10% (R$)'].sum())
    arquivo = excel(resumo, trechos, agenda_df, pontos, {'Teste': '=1+1'})
    wb = load_workbook(BytesIO(arquivo))
    assert len(wb.sheetnames) == 5
    assert wb['Premissas']['B2'].data_type == 's'
    wb.close()
    relatorio = html(resumo, trechos, agenda_df, pontos, {'Preço': 6.34}).decode('utf-8')
    assert 'NISSAN B.LEME' not in relatorio
    assert 'window.print()' in relatorio
